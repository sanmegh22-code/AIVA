from openpyxl import Workbook
from openpyxl.styles import (
    Font,
    PatternFill,
    Border,
    Side,
    Alignment,
)
from openpyxl.utils import get_column_letter

from app.services.reports.analytics import ReportAnalytics


class ExcelGenerator:

    HEADER_FILL = PatternFill(
        start_color="1F2937",
        end_color="1F2937",
        fill_type="solid",
    )

    SUBHEADER_FILL = PatternFill(
        start_color="2563EB",
        end_color="2563EB",
        fill_type="solid",
    )

    WHITE_FONT = Font(
        bold=True,
        color="FFFFFF",
    )

    TITLE_FONT = Font(
        bold=True,
        size=18,
    )

    BOLD_FONT = Font(
        bold=True,
    )

    BORDER = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    CENTER = Alignment(
        horizontal="center",
        vertical="center",
    )

    @staticmethod
    def build_excel(db):

        report = ReportAnalytics.get_complete_report(db)

        workbook = Workbook()

        ExcelGenerator.summary_sheet(
            workbook,
            report,
        )

        ExcelGenerator.products_sheet(
            workbook,
            report,
        )

        ExcelGenerator.analytics_sheet(
            workbook,
            report,
        )

        ExcelGenerator.auto_resize(workbook)

        return workbook

    @staticmethod
    def summary_sheet(workbook, report):

        ws = workbook.active

        ws.title = "Summary"

        ws.merge_cells("A1:B1")

        cell = ws["A1"]

        cell.value = "AIVA Inventory Report"

        cell.font = ExcelGenerator.TITLE_FONT

        ws.append([])

        ws.append(["Metric", "Value"])

        for c in ws[3]:
            c.fill = ExcelGenerator.HEADER_FILL
            c.font = ExcelGenerator.WHITE_FONT
            c.border = ExcelGenerator.BORDER
            c.alignment = ExcelGenerator.CENTER

        summary = report["summary"]

        rows = [

            ("Products", summary["total_products"]),

            ("Categories", summary["total_categories"]),

            ("Suppliers", summary["total_suppliers"]),

            ("Warehouses", summary["total_warehouses"]),

            ("Inventory", summary["total_inventory"]),

            (
                "Inventory Value",
                summary["inventory_value"],
            ),

            (
                "Average Inventory Value",
                summary["average_inventory_value"],
            ),

            (
                "Average Quantity",
                summary["average_quantity"],
            ),

            (
                "Stock Health %",
                summary["stock_health"],
            ),

            (
                "Warehouse Utilization",
                summary["warehouse_utilization"],
            ),

            (
                "Low Stock",
                summary["low_stock"],
            ),

            (
                "Out Of Stock",
                summary["out_of_stock"],
            ),
        ]

        for row in rows:

            ws.append(row)

            for cell in ws[ws.max_row]:

                cell.border = ExcelGenerator.BORDER

    @staticmethod
    def products_sheet(workbook, report):

        ws = workbook.create_sheet("Top Products")

        ws.append([
            "Product",
            "SKU",
            "Price",
            "Quantity",
            "Category",
        ])

        for cell in ws[1]:

            cell.fill = ExcelGenerator.SUBHEADER_FILL

            cell.font = ExcelGenerator.WHITE_FONT

            cell.border = ExcelGenerator.BORDER

        for product in report["top_products"]:

            ws.append([

                product.name,

                product.sku,

                product.price,

                product.total_quantity,

                product.category,

            ])

            for cell in ws[ws.max_row]:

                cell.border = ExcelGenerator.BORDER

    @staticmethod
    def analytics_sheet(workbook, report):

        ws = workbook.create_sheet("Analytics")

        ws.append([
            "Analytics",
            "Value",
        ])

        for cell in ws[1]:

            cell.fill = ExcelGenerator.SUBHEADER_FILL

            cell.font = ExcelGenerator.WHITE_FONT

            cell.border = ExcelGenerator.BORDER

        summary = report["summary"]

        movement = report["movements"]

        analytics = [

            (
                "Inventory Value",
                summary["inventory_value"],
            ),

            (
                "Average Product Value",
                summary["average_inventory_value"],
            ),

            (
                "Average Quantity",
                summary["average_quantity"],
            ),

            (
                "Stock Health",
                summary["stock_health"],
            ),

            (
                "Warehouse Utilization",
                summary["warehouse_utilization"],
            ),

            (
                "Stock In",
                movement["total_in"],
            ),

            (
                "Stock Out",
                movement["total_out"],
            ),

            (
                "Net Stock",
                movement["net_stock"],
            ),

        ]

        for row in analytics:

            ws.append(row)

            for cell in ws[ws.max_row]:

                cell.border = ExcelGenerator.BORDER

    @staticmethod
    def auto_resize(workbook):

        for sheet in workbook.worksheets:

            for column in sheet.columns:

                length = 0

                column_letter = get_column_letter(
                    column[0].column
                )

                for cell in column:

                    try:

                        if len(str(cell.value)) > length:

                            length = len(str(cell.value))

                    except Exception:

                        pass

                sheet.column_dimensions[
                    column_letter
                ].width = length + 5